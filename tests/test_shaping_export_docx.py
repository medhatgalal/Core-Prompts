"""Native DOCX regressions. Run with the approved document runtime.

CI installs optional document dependencies. The ordinary stdlib-only environment
may skip this module; skipped cases are never counted as document proof.
"""
import importlib.util
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

AVAILABLE = all(importlib.util.find_spec(m) is not None for m in ('docx', 'PIL'))
SOURCE = Path(__file__).resolve().parents[1] / 'sources/capability-resources/engos-delivery-artifact-embed/scripts/artifact_export.py'
SPEC = importlib.util.spec_from_file_location('export_docx_under_test', SOURCE)
export = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(export)


@unittest.skipUnless(AVAILABLE, 'requires approved python-docx and Pillow runtime')
class DocxRegressions(unittest.TestCase):
    def setUp(self):
        from PIL import Image
        self.temp = tempfile.TemporaryDirectory(prefix='shaping-docx-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.bundle = self.root/'bundle'; self.bundle.mkdir()
        self.images = self.root/'images'; self.images.mkdir()
        (self.bundle/'pitch.md').write_text('# Test pitch\n\n```sh\n> output.txt\necho `whoami`\n```\n')
        for name in ('contracts.md','security-owners.md'):
            (self.bundle/name).write_text('| ID | Meaning |\n| --- | --- |\n| A | preserved |\n')
        for name in export.DIAGRAMS:
            (self.bundle/name).write_text('flowchart LR\n A-->B\n')
        self.data = export.bundle(self.bundle)
        manifest = {}
        for d in self.data['diagrams']:
            filename = d['id'].replace('.mmd','.png')
            stream = io.BytesIO(); Image.new('RGB',(50,30),'red').save(stream,format='PNG')
            raw = stream.getvalue(); (self.images/filename).write_bytes(raw)
            manifest[d['id']] = {'source_sha256':d['sha256'],'png':filename,'png_sha256':export.digest(raw)}
        (self.images/'render-manifest.json').write_text(json.dumps(manifest))
        self.output = self.root/'pitch.docx'

    def test_fenced_code_preserves_literal_operators(self):
        from docx import Document
        export.docx_document(self.data,self.images,self.output)
        text = '\n'.join(p.text for p in Document(self.output).paragraphs)
        self.assertIn('> output.txt\necho `whoami`',text)

    def test_concurrent_creation_is_not_overwritten(self):
        from docx.document import Document
        real_save = Document.save
        def race(document, target):
            self.output.write_bytes(b'user replacement')
            return real_save(document,target)
        with patch.object(Document,'save',race):
            with self.assertRaises(export.ExportError):
                export.docx_document(self.data,self.images,self.output)
        self.assertEqual(self.output.read_bytes(),b'user replacement')

    def test_only_verified_image_bytes_are_embedded(self):
        import docx
        from PIL import Image
        original = (self.images/'component.png').read_bytes()
        real_document = docx.Document
        def mutate(*args,**kwargs):
            Image.new('RGB',(50,30),'blue').save(self.images/'component.png')
            return real_document(*args,**kwargs)
        with patch.object(docx,'Document',mutate):
            export.docx_document(self.data,self.images,self.output)
        saved = real_document(self.output)
        image = saved.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
        self.assertEqual(saved.part.related_parts[image].blob,original)

    def test_framed_native_styles_lists_and_links_without_images(self):
        from docx import Document
        (self.bundle/'pitch.md').write_text('# Problem\n\nA **clear** problem\nwith *context* and [evidence](https://example.org).\n\n## Why now\n\n- One\n- Two\n\n### Bounds\n\n3. Third\n4. Fourth\n')
        data = export.bundle(self.bundle,profile='framed')
        export.docx_document(data,None,self.output)
        saved = Document(self.output)
        self.assertEqual(len(saved.inline_shapes),0)
        self.assertEqual(saved.sections[0].page_width.inches,8.5)
        self.assertEqual(saved.styles['Title'].font.size.pt,18)
        self.assertEqual(str(saved.styles['Title'].font.color.rgb),'1B5BBB')
        self.assertEqual(str(saved.styles['Heading 2'].font.color.rgb),'157A3B')
        self.assertEqual(str(saved.styles['Caption'].font.color.rgb),'495565')
        self.assertIsNone(saved.styles['Caption'].font.color.theme_color)
        self.assertEqual(len([p for p in saved.paragraphs if p.style.name=='List Bullet']),2)
        self.assertTrue(saved.element.xpath('//w:hyperlink'))
        self.assertIn('https://example.org',[r.target_ref for r in saved.part.rels.values()])
        self.assertIn('problem with', saved.paragraphs[1].text)
        self.assertNotIn('**',saved.paragraphs[1].text)
        self.assertTrue(saved.element.xpath('//w:numPr'))

    def test_shaped_native_compact_tables_all_fields_and_anchored_figures(self):
        from docx import Document
        (self.bundle/'pitch.md').write_text('# Proposal\n\n## Architecture\n\nContext A\n\n## Sequence\n\nContext B\n\n## Contracts\n\nContext C\n')
        headers = ['ID','Interface','Purpose','State','Owner','Producer','Consumer','Inputs','Outputs','Errors','Retry','Consistency','Security']
        values = ['C-1'] + ['detail '+h for h in headers[1:]]
        (self.bundle/'contracts.md').write_text('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*13)+' |\n| '+' | '.join(values)+' |\n')
        (self.bundle/'security-owners.md').write_text('| Responsibility | Owner | How enforced | Evidence |\n| --- | --- | --- | --- |\n| Access | Team | Policy | Pending |\n')
        mapping = {'schema_version':'ShapingPresentation.v1','diagrams':{}}
        manifest = json.loads((self.images/'render-manifest.json').read_text())
        for name,section in zip(export.DIAGRAMS,['architecture','sequence','contracts']):
            item = {'section_id':'pitch.md#'+section,'caption':name+' caption','alt':name+' alt'}
            mapping['diagrams'][name] = item
            svg = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 50 30"><text x="1" y="20">Fixture</text></svg>'
            filename = name.replace('.mmd','.svg'); (self.images/filename).write_bytes(svg)
            manifest[name].update(item,svg=filename,svg_sha256=export.digest(svg),svg_dimensions=[50,30],png_dimensions=[50,30],renderer={'version':'fixture 1','config':{'htmlLabels':False,'securityLevel':'strict'}})
        (self.bundle/'presentation.json').write_text(json.dumps(mapping))
        (self.images/'render-manifest.json').write_text(json.dumps(manifest))
        data = export.bundle(self.bundle,profile='shaped',presentation='presentation.json')
        export.docx_document(data,self.images,self.output)
        saved = Document(self.output)
        self.assertEqual(len(saved.inline_shapes),3)
        self.assertEqual(len(saved.tables[0].columns),5)
        all_cells = [c.text for t in saved.tables for r in t.rows for c in r.cells]
        for value in values: self.assertIn(value,all_cells)
        self.assertTrue(saved.element.xpath('//w:bookmarkStart'))
        self.assertTrue(saved.element.xpath('//w:hyperlink[@w:anchor]'))
        self.assertTrue(saved.element.xpath('//w:tblHeader'))
        body = saved.element.xml
        self.assertLess(body.index('component.mmd caption'),body.index('>Sequence</w:t>'))
        self.assertIn('component.mmd alt',body)

    def test_local_fragments_and_colliding_headings_resolve(self):
        from docx import Document
        from docx.oxml.ns import qn
        (self.bundle/'framed.md').write_text('# Bounds\n\n[See bounds](#bounds)\n\n## Bounds\n\n## Bounds-2\n')
        export.docx_document(export.bundle(self.bundle,profile='framed'),None,self.output)
        saved = Document(self.output)
        names = {el.get(qn('w:name')) for el in saved.element.xpath('//w:bookmarkStart')}
        targets = [el.get(qn('w:anchor')) for el in saved.element.xpath('//w:hyperlink')]
        self.assertEqual(len(names),4)  # Document destination plus three distinct headings.
        self.assertTrue(targets and all(t in names for t in targets))

    def test_tall_legacy_figures_fit_page_with_caption(self):
        from docx import Document
        from PIL import Image
        manifest = json.loads((self.images/'render-manifest.json').read_text())
        for diagram in self.data['diagrams']:
            path = self.images/manifest[diagram['id']]['png']
            Image.new('RGB',(100,200),'red').save(path)
            manifest[diagram['id']]['png_sha256'] = export.digest(path.read_bytes())
        (self.images/'render-manifest.json').write_text(json.dumps(manifest))
        export.docx_document(self.data,self.images,self.output)
        saved = Document(self.output)
        last = saved.sections[-1]
        self.assertGreater(last.page_height,last.page_width)
        usable = (last.page_height-last.top_margin-last.bottom_margin)/914400
        self.assertTrue(all(s.height.inches <= usable-.5 for s in saved.inline_shapes))

    def test_quote_and_relative_links_are_readable_and_bound(self):
        from docx import Document
        from docx.oxml.ns import qn
        (self.bundle/'framed.md').write_text('# Frame\n\n> Preserve **uncertainty**\n> and context.\n\n[Research](research.md#bounds) [Document](research.md) [Source](../source.md)\n')
        (self.bundle/'research.md').write_text('# Research\n\n## Bounds\n\nUnknown.\n')
        data = export.bundle(self.bundle,['framed.md','research.md'],profile='framed')
        export.docx_document(data,None,self.output)
        saved = Document(self.output)
        quote = next(p for p in saved.paragraphs if p.style.name=='Quote')
        self.assertEqual(quote.text,'Preserve uncertainty and context.')
        names = {e.get(qn('w:name')) for e in saved.element.xpath('//w:bookmarkStart')}
        self.assertTrue(all(e.get(qn('w:anchor')) in names for e in saved.element.xpath('//w:hyperlink')))
        self.assertIn('Source (../source.md)',saved.element.xml)

    def test_detail_semantics_pagination_and_record_header(self):
        from docx import Document
        raw = 'Use `fieldName` and **keep** unknown values; [evidence](https://example.org).'
        table = self.data['documents'][1]['blocks'][0]
        table['presentation'] = {'headers':['ID','Purpose'],'records':[{'id':'contracts.md--C4','key':'C4','summary':['C4','Keep values'],'fields':[('ID','C4'),('Purpose','Keep values'),('Meaning',raw)]}]}
        self.data['documents'][0]['blocks'].append({'kind':'paragraph','text':'[Contract C4](#contracts.md--C4)','document_id':'pitch.md','section_id':'pitch.md#test-pitch'})
        export.docx_document(self.data,self.images,self.output)
        saved = Document(self.output)
        detail = next(t for t in saved.tables if t.cell(0,0).text=='Field')
        self.assertEqual(detail.cell(0,1).text,'C4')
        self.assertEqual(detail.cell(3,1).text,'Use fieldName and keep unknown values; evidence.')
        self.assertEqual(export.semantic_text(raw,'contracts.md',self.data),detail.cell(3,1).text)
        self.assertTrue(any(r.text=='fieldName' and r.font.name=='Courier New' for r in detail.cell(3,1).paragraphs[0].runs))
        for row in detail.rows: self.assertTrue(row._tr.xpath('./w:trPr/w:cantSplit'))
        self.assertTrue(detail.rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
        self.assertTrue(detail.cell(1,1).paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(detail.cell(2,1).paragraphs[0].paragraph_format.keep_with_next)
        self.assertEqual(table['presentation']['records'][0]['fields'][-1][1],raw)
        from docx.oxml.ns import qn
        names = {e.get(qn('w:name')) for e in saved.element.xpath('//w:bookmarkStart')}
        self.assertTrue(all(e.get(qn('w:anchor')) in names for e in saved.element.xpath('//w:hyperlink[@w:anchor]')))

    def test_oversized_native_row_fails_without_publishing(self):
        table = self.data['documents'][1]['blocks'][0]
        table['rows'][0][1] = 'Oversized evidence sentence. '*2000
        with self.assertRaisesRegex(export.ExportError,'row.*page'):
            export.docx_document(self.data,self.images,self.output)
        self.assertFalse(self.output.exists())

    def test_figure_title_caption_and_legend_stay_together(self):
        from docx import Document
        self.data['diagrams'][0].update(title='Supplied title',legend='Dashed means proposed; solid means existing.',caption='Supplied caption')
        manifest = json.loads((self.images/'render-manifest.json').read_text())
        manifest['component.mmd'].update(title='Supplied title',legend='Dashed means proposed; solid means existing.',caption='Supplied caption')
        manifest['component.mmd']['presentation_sha256'] = export.diagram_metadata_hash(self.data['diagrams'][0])
        (self.images/'render-manifest.json').write_text(json.dumps(manifest))
        export.docx_document(self.data,self.images,self.output)
        saved = Document(self.output)
        paragraphs = saved.paragraphs
        title = next(i for i,p in enumerate(paragraphs) if p.text=='Supplied title')
        self.assertEqual(paragraphs[title+1].text,'Supplied caption')
        self.assertTrue(paragraphs[title+2]._p.xpath('.//w:drawing'))
        self.assertEqual(paragraphs[title+3].text,'Dashed means proposed; solid means existing.')
        self.assertTrue(all(paragraphs[i].paragraph_format.keep_with_next for i in range(title,title+3)))
        self.assertFalse(paragraphs[title+3].paragraph_format.keep_with_next)
        self.assertTrue(all(paragraphs[i].paragraph_format.keep_together for i in range(title,title+4)))

    def test_figure_plain_metadata_is_literal_in_docx(self):
        from docx import Document
        fields={'title':'A * B * C','caption':'Literal **caption** value','legend':'Literal `node` label'}
        self.data['diagrams'][0].update(fields)
        manifest=json.loads((self.images/'render-manifest.json').read_text())
        manifest['component.mmd'].update(fields)
        manifest['component.mmd']['presentation_sha256']=export.diagram_metadata_hash(self.data['diagrams'][0])
        (self.images/'render-manifest.json').write_text(json.dumps(manifest))
        export.docx_document(self.data,self.images,self.output)
        text=[p.text for p in Document(self.output).paragraphs]
        for value in fields.values(): self.assertIn(value,text)


if __name__ == '__main__':
    unittest.main()
