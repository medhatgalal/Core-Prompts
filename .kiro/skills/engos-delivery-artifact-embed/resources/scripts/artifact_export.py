"""Local shaping exports. Never publishes or certifies review/visual quality.

Supported Markdown subset: headings, paragraphs/lists, fenced code, pipe tables.
Original UTF-8 Markdown and Mermaid are preserved verbatim in JSON. Complex table
syntax fails explicitly instead of silently dropping rows. DOCX is optional.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import hashlib
import html
import io
import json
import os
import re
import selectors
import signal
import stat
import subprocess
import sys
import tempfile
import time
from pathlib import Path

DOCUMENTS = ("pitch.md", "contracts.md", "security-owners.md")
DIAGRAMS = ("component.mmd", "sequence.mmd", "data-flow.mmd")


class ExportError(ValueError):
    pass


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_source(root: Path, name: str) -> bytes:
    rel = Path(name)
    if rel.is_absolute() or ".." in rel.parts or not rel.parts:
        raise ExportError(f"unsafe source path: {name}")
    current = root
    for part in rel.parts:
        current = current / part
        if current.is_symlink():
            raise ExportError(f"symlink source refused: {name}")
    if not current.is_file() or not current.resolve().is_relative_to(root.resolve()):
        raise ExportError(f"missing or outside source: {name}")
    return current.read_bytes()


def cells(line: str) -> list[str]:
    # Escaped pipe is content, never an extra cell. Preserve other backslashes.
    return [x.strip().replace(r"\|", "|") for x in re.split(r"(?<!\\)\|", line.strip()[1:-1])]


def blocks(text: str) -> list[dict]:
    lines, output, i = text.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            language, content = line[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                content.append(lines[i]); i += 1
            if i == len(lines):
                raise ExportError("unclosed code fence")
            output.append({"kind": "code", "language": language, "text": "\n".join(content)})
            i += 1
        elif line.strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not lines[i].strip().endswith("|"):
                    raise ExportError("tables must have explicit closing pipes")
                rows.append(cells(lines[i])); i += 1
            if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
                raise ExportError("pipe table requires header and separator")
            if any(len(row) != len(rows[0]) for row in rows):
                raise ExportError("inconsistent table column count")
            output.append({"kind": "table", "headers": rows[0], "rows": rows[2:]})
        elif re.match(r"#{1,6} ", line):
            level = len(line) - len(line.lstrip("#"))
            output.append({"kind": "heading", "level": level, "text": line[level:].strip()})
            i += 1
        else:
            output.append({"kind": "paragraph", "text": line})
            i += 1
    return output


def bundle(root: Path, documents=DOCUMENTS) -> dict:
    if root.is_symlink() or not root.is_dir():
        raise ExportError("bundle root must be a real directory")
    if len(set(documents)) != len(documents) or not set(DOCUMENTS).issubset(documents):
        raise ExportError("include each required document once")
    result = {"schema_version": "ShapingDocumentExport.v1", "documents": [], "diagrams": []}
    inventory = {}
    for name in (*documents, *DIAGRAMS):
        raw = read_source(root, name)
        content = raw.decode("utf-8")
        if not content.strip():
            raise ExportError(f"empty required source: {name}")
        sha = digest(raw)
        inventory[name] = sha
        if name in DIAGRAMS:
            result["diagrams"].append({"id": name, "sha256": sha, "mermaid": content})
        else:
            parsed = blocks(content)
            section, seen = name + '#preamble', {}
            for block in parsed:
                if block['kind'] == 'heading':
                    slug = re.sub(r'[^a-z0-9]+', '-', block['text'].lower()).strip('-') or 'section'
                    seen[slug] = seen.get(slug, 0) + 1
                    section = name + '#' + slug + (f'-{seen[slug]}' if seen[slug] > 1 else '')
                    block['id'] = section
                block['section_id'] = section
            if name in ("contracts.md", "security-owners.md") and not any(
                b["kind"] == "table" and b["rows"] for b in parsed
            ):
                raise ExportError(f"required table has no data rows: {name}")
            result["documents"].append({"id": name, "sha256": sha, "markdown": content, "blocks": parsed})
    result["source_inventory"] = inventory
    result["source_hash"] = digest(json.dumps(inventory, sort_keys=True).encode())
    result["table_row_counts"] = {
        doc["id"]: [len(b["rows"]) for b in doc["blocks"] if b["kind"] == "table"]
        for doc in result["documents"]
    }
    result["verification"] = "source_export_only; semantic, visual and saved-target review required"
    return result


def html_document(data: dict) -> str:
    out = ['<!doctype html><html lang="en"><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           '<title>Shaped pitch</title><style>body{font:16px/1.5 system-ui;max-width:1100px;margin:auto;padding:32px;color:#162b40}table{border-collapse:collapse;width:100%}th,td{border:1px solid #d9d9d9;padding:8px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#e3f2fd}pre{white-space:pre-wrap}.table,.mermaid{overflow:auto}h1,h2,h3{color:#183b56}</style><main>']
    esc = html.escape
    for doc in data["documents"]:
        out.append(f'<section data-source="{esc(doc["id"], quote=True)}">')
        for b in doc["blocks"]:
            if b["kind"] == "heading":
                n = b["level"]
                out.append(f'<h{n}>{esc(b["text"])}</h{n}>')
            elif b["kind"] == "table":
                out.append('<div class="table"><table><thead><tr>' + ''.join('<th>'+esc(c)+'</th>' for c in b["headers"]) + '</tr></thead><tbody>')
                for row in b["rows"]:
                    out.append('<tr>' + ''.join('<td>'+esc(c)+'</td>' for c in row) + '</tr>')
                out.append('</tbody></table></div>')
            elif b["kind"] == "code":
                out.append('<pre>'+esc(b["text"])+'</pre>')
            else:
                out.append('<p>'+esc(b["text"])+'</p>')
        out.append('</section>')
    for diagram in data["diagrams"]:
        out.append('<h2>'+esc(diagram["id"].replace('.mmd', '').replace('-', ' ').title())+'</h2>')
        out.append('<pre class="mermaid">'+esc(diagram["mermaid"])+'</pre>')
    out.append('<p>Source revision '+data["source_hash"]+'</p>')
    out.append('<script type="module">import mermaid from "https://cdn.jsdelivr.net/npm/mermaid@11.17.2/dist/mermaid.esm.min.mjs";mermaid.initialize({startOnLoad:false,securityLevel:"strict"});await mermaid.run({querySelector:".mermaid"});for(const svg of document.querySelectorAll(".mermaid svg")){svg.style.width=svg.viewBox.baseVal.width+"px";svg.style.maxWidth="none";}</script></main></html>')
    return '\n'.join(out)


@contextmanager
def directory_handle(path: Path):
    """Pin each physical output directory; never traverse a replaced symlink.

    POSIX dir-fd/no-follow primitives are required rather than silently falling
    back to a racy pathname publication on unsupported hosts.
    """
    if os.name != 'posix' or not hasattr(os, 'O_NOFOLLOW'):
        raise ExportError('safe output publication requires POSIX directory descriptors')
    path = Path(os.path.abspath(path))
    fd = os.open(path.anchor, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        for part in path.parts[1:]:
            try:
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=fd)
            except FileNotFoundError:
                try:
                    os.mkdir(part, mode=0o700, dir_fd=fd)
                except FileExistsError:
                    pass
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=fd)
            os.close(fd); fd = child
        yield fd
    except OSError as error:
        raise ExportError(f'unsafe or unavailable output directory: {error}') from error
    finally:
        os.close(fd)


def write_at(directory: int, name: str, raw: bytes, identical=True) -> None:
    if not name or Path(name).name != name:
        raise ExportError('invalid output filename')
    try:
        fd = os.open(name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                     0o600, dir_fd=directory)
    except FileExistsError:
        if identical:
            try:
                fd = os.open(name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
                with os.fdopen(fd, 'rb') as stream:
                    if stat.S_ISREG(os.fstat(stream.fileno()).st_mode) and stream.read(len(raw)+1) == raw:
                        return
            except OSError:
                pass
        raise ExportError(f'output exists; use a new revision directory: {name}')
    with os.fdopen(fd, 'wb') as stream:
        stream.write(raw); stream.flush(); os.fsync(stream.fileno())
    os.fsync(directory)


def write_new(path: Path, raw: bytes) -> None:
    with directory_handle(path.parent) as directory:
        write_at(directory, path.name, raw)


def run_renderer(argv: list[str], timeout: float) -> None:
    """Finite wall-clock limit, bounded diagnostics, owned process-group cleanup."""
    if timeout <= 0 or timeout > 300 or os.name != 'posix':
        raise ExportError('renderer needs POSIX and a timeout in (0, 300] seconds')
    diagnostic = bytearray()
    process = subprocess.Popen(argv, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                               stdin=subprocess.DEVNULL, start_new_session=True)
    deadline = time.monotonic() + timeout
    timed_out = False
    try:
        with selectors.DefaultSelector() as selector:
            selector.register(process.stdout, selectors.EVENT_READ)
            while selector.get_map():
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    timed_out = True
                    break
                for key, _ in selector.select(min(remaining, .1)):
                    chunk = os.read(key.fileobj.fileno(), 8192)
                    if not chunk:
                        selector.unregister(key.fileobj)
                    elif len(diagnostic) < 16384:
                        diagnostic.extend(chunk[:16384-len(diagnostic)])
            if not timed_out:
                try:
                    process.wait(timeout=max(.001, deadline-time.monotonic()))
                except subprocess.TimeoutExpired:
                    timed_out = True
    finally:
        # Only this explicitly started group, never a broad renderer/process match.
        for sig in (signal.SIGTERM, signal.SIGKILL):
            try:
                os.killpg(process.pid, sig)
            except ProcessLookupError:
                pass
            if sig == signal.SIGTERM:
                try:
                    process.wait(timeout=.5)
                except subprocess.TimeoutExpired:
                    pass
        process.wait()
        process.stdout.close()
    if timed_out:
        raise ExportError(f'renderer timed out after {timeout} seconds')
    if process.returncode:
        raise ExportError(f'renderer exit {process.returncode}: {diagnostic.decode("utf-8",errors="replace")}')


def render(data: dict, root: Path, output: Path, executable: str, timeout=60) -> dict:
    # Refuse symlink parents before creating a directory or invoking a renderer.
    with directory_handle(output.parent) as parent:
        try:
            os.mkdir(output.name, mode=0o700, dir_fd=parent)
        except FileExistsError as error:
            raise ExportError('render to a fresh directory') from error
        directory = os.open(output.name, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=parent)
        try:
            receipts = {}
            with tempfile.TemporaryDirectory(prefix='engos-mermaid-') as temporary:
                scratch = Path(temporary).resolve()
                for diagram in data['diagrams']:
                    source = scratch / diagram['id']
                    source.write_text(diagram['mermaid'], encoding='utf-8')
                    png = scratch / (source.stem + '.png')
                    run_renderer([executable, '-i', str(source), '-o', str(png), '-b', 'white', '-w', '1400', '-s', '2'], timeout)
                    if digest(read_source(root, diagram['id'])) != diagram['sha256']:
                        raise ExportError('source changed during rendering')
                    raw = read_source(scratch, png.name)
                    write_at(directory, png.name, raw, identical=False)
                    receipts[diagram['id']] = {'source_sha256':diagram['sha256'], 'png':png.name,
                                               'png_sha256':digest(raw), 'pixel_inspection':'pending'}
            write_at(directory, 'render-manifest.json', json.dumps(receipts, indent=2).encode(), identical=False)
            return receipts
        finally:
            os.close(directory)


def docx_document(data: dict, images: Path, output: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.section import WD_SECTION_START
    from PIL import Image

    if output.exists() or any(p.is_symlink() for p in (output, *output.parents)):
        raise ExportError("DOCX output exists; use a fresh revision")
    receipts = json.loads(read_source(images, 'render-manifest.json'))
    verified_images = {}
    # Check every image/source binding BEFORE constructing a document.
    for diagram in data["diagrams"]:
        entry = receipts.get(diagram["id"], {})
        if entry.get('source_sha256') != diagram['sha256']:
            raise ExportError("stale or missing render source binding")
        png = read_source(images, entry.get('png', ''))
        if digest(png) != entry.get('png_sha256'):
            raise ExportError("render image hash mismatch")
        verified_images[diagram['id']] = png
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    for name, size in [('Normal',11),('Title',18),('Heading 1',13),('Heading 2',11.5)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor.from_string('183B56' if name.startswith('Heading') else '000000')
        style.font.bold = name != 'Normal'
        style.paragraph_format.space_after = Pt(6)
    title_added = False

    def paragraph(text, style=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.widow_control = True
        p.paragraph_format.keep_together = len(text.split()) < 120
        # Preserve link destinations as visible references without adding remote
        # content. Inline code/bold become formatting; JSON keeps source bytes.
        text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
        if text.startswith('> '):
            text = text[2:]
        for chunk in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if chunk.startswith('**') and chunk.endswith('**'):
                p.add_run(chunk[2:-2]).bold = True
            elif chunk.startswith('`') and chunk.endswith('`'):
                p.add_run(chunk[1:-1]).font.name = 'Courier New'
            else:
                p.add_run(chunk)
        return p

    for item in data['documents']:
        for block in item['blocks']:
            if block['kind'] == 'heading':
                style = 'Title' if not title_added else ('Heading 1' if block['level'] <= 2 else 'Heading 2')
                paragraph(block['text'], style); title_added = True
            elif block['kind'] == 'table':
                table = doc.add_table(rows=0, cols=len(block['headers']))
                table.style = 'Table Grid'
                table.autofit = False
                if len(block['headers']) == 6 and block['headers'][0].strip().lower() == 'id':
                    weights = [5, 16, 12, 15, 30, 22]
                elif len(block['headers']) == 3:
                    weights = [24, 34, 42]
                else:
                    weights = [1] * len(block['headers'])
                widths = [Inches(10.4 * w / sum(weights)) for w in weights]
                for column, width in zip(table.columns, widths):
                    column.width = width
                for index, values in enumerate([block['headers'], *block['rows']]):
                    row = table.add_row()
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell, value, width in zip(row.cells, values, widths):
                        cell.width = width
                        cell.text = value
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(4)
                            p.paragraph_format.keep_with_next = index == 0
                            for r in p.runs:
                                r.font.size = Pt(10); r.bold = index == 0
                        borders = OxmlElement('w:tcBorders')
                        for edge in ('top','left','bottom','right'):
                            el = OxmlElement('w:'+edge)
                            for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                                el.set(qn('w:'+key), val)
                            borders.append(el)
                        cell._tc.get_or_add_tcPr().append(borders)
                        if index == 0:
                            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E3F2FD')
                            cell._tc.get_or_add_tcPr().append(shade)
                    if index == 0:
                        row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                doc.add_paragraph()
            elif block['kind'] == 'code':
                p = doc.add_paragraph()
                p.add_run(block['text']).font.name = 'Courier New'
            else:
                paragraph(block['text'])
    for diagram in data['diagrams']:
        png = verified_images[diagram['id']]
        with Image.open(io.BytesIO(png)) as im:
            portrait = im.height > .75 * im.width
            if portrait != (section.page_height > section.page_width):
                section = doc.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width = Inches(8.3 if portrait else 11.7)
                section.page_height = Inches(11.7 if portrait else 8.3)
            width = min(6.8 if portrait else 10.0, (9.0 if portrait else 5.8) * im.width / im.height)
        doc.add_paragraph(diagram['id'].replace('.mmd','').replace('-',' ').title(), 'Heading 1')
        doc.add_picture(io.BytesIO(png), width=Inches(width))
    doc.add_paragraph('Source revision '+data['source_hash'])
    stream = io.BytesIO()
    doc.save(stream)
    with directory_handle(output.parent) as parent:
        write_at(parent, output.name, stream.getvalue(), identical=False)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('action', choices=('json','html','render','docx'))
    p.add_argument('--bundle', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    p.add_argument('--documents', nargs='+', default=list(DOCUMENTS))
    p.add_argument('--mmdc', default='mmdc', help='existing trusted Mermaid CLI executable; no installation performed')
    p.add_argument('--images', type=Path)
    p.add_argument('--timeout-seconds', type=float, default=60, help='renderer wall-clock limit per diagram, maximum 300')
    args = p.parse_args(argv)
    try:
        data = bundle(args.bundle, args.documents)
        if args.action == 'json':
            write_new(args.output, (json.dumps(data, indent=2, ensure_ascii=False)+'\n').encode())
        elif args.action == 'html':
            write_new(args.output, html_document(data).encode())
        elif args.action == 'render':
            render(data, args.bundle, args.output, args.mmdc, args.timeout_seconds)
        elif args.images is None:
            raise ExportError('docx requires --images with source-bound render-manifest.json')
        else:
            docx_document(data, args.images, args.output)
        print(json.dumps({'status':'created_or_identical','source_hash':data['source_hash'],
                          'table_row_counts':data['table_row_counts'], 'diagrams':len(data['diagrams']),
                          'visual_review':'pending','saved_target_review':'not_performed'}))
        return 0
    except (ExportError, OSError, UnicodeError, json.JSONDecodeError, subprocess.CalledProcessError, ImportError) as error:
        print(json.dumps({'status':'error','detail':str(error)}), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
