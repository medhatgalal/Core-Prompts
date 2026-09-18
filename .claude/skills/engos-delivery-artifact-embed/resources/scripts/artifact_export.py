"""Local shaping exports. Never publishes or certifies review/visual quality.

Supported Markdown subset: headings, paragraphs/lists, fenced code, pipe tables.
Original UTF-8 Markdown and Mermaid are preserved verbatim in JSON. Complex table
syntax fails explicitly instead of silently dropping rows. DOCX is optional.
"""
from __future__ import annotations

import argparse
import base64
from contextlib import contextmanager
import hashlib
import html
import io
import json
import math
import os
import posixpath
import re
import selectors
import signal
import stat
import struct
import subprocess
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
import zlib
from pathlib import Path
from urllib.parse import urlsplit

DOCUMENTS = ("pitch.md", "contracts.md", "security-owners.md")
DIAGRAMS = ("component.mmd", "sequence.mmd", "data-flow.mmd")
BLUE, GREEN = '1B5BBB', '157A3B'
RENDER_CONFIG = {'securityLevel': 'strict', 'htmlLabels': False,
                 'flowchart': {'htmlLabels': False}}
FIGURE_FIELDS = ('section_id','caption','alt','title','legend')


class ExportError(ValueError):
    pass


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_source(root: Path, name: str) -> bytes:
    if not isinstance(name, str):
        raise ExportError('source path must be relative to a real directory')
    rel = Path(name)
    if rel.is_absolute() or ".." in rel.parts or not rel.parts:
        raise ExportError(f"unsafe source path: {name}")
    try:
        with directory_handle(root / rel.parent, create=False) as directory:
            fd = os.open(rel.name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
            with os.fdopen(fd, 'rb') as source:
                before = os.fstat(source.fileno())
                if not stat.S_ISREG(before.st_mode):
                    raise ExportError('source must be a regular file: '+name)
                raw = source.read()
                after = os.fstat(source.fileno())
                if (before.st_size,before.st_mtime_ns,before.st_ctime_ns) != (after.st_size,after.st_mtime_ns,after.st_ctime_ns):
                    raise ExportError('source changed while reading: '+name)
                return raw
    except OSError as error:
        raise ExportError('unsafe or unavailable source: '+name) from error


def cells(line: str) -> list[str]:
    # Escaped pipe is content, never an extra cell. Preserve other backslashes.
    return [x.strip().replace(r"\|", "|") for x in re.split(r"(?<!\\)\|", line.strip()[1:-1])]


def blocks(text: str) -> list[dict]:
    lines, output, i = text.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            language, content = line[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                content.append(lines[i]); i += 1
            if i == len(lines):
                raise ExportError("unclosed code fence")
            output.append({"kind": "code", "language": language, "text": "\n".join(content)})
            i += 1
        elif line.strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not lines[i].strip().endswith("|"):
                    raise ExportError("tables must have explicit closing pipes")
                rows.append(cells(lines[i])); i += 1
            if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
                raise ExportError("pipe table requires header and separator")
            if any(len(row) != len(rows[0]) for row in rows):
                raise ExportError("inconsistent table column count")
            output.append({"kind": "table", "headers": rows[0], "rows": rows[2:]})
        elif line.startswith('> '):
            content = []
            while i < len(lines) and lines[i].startswith('> '):
                content.append(lines[i][2:]); i += 1
            output.append({'kind':'quote','text':' '.join(content)})
        elif re.match(r"#{1,6} ", line):
            level = len(line) - len(line.lstrip("#"))
            output.append({"kind": "heading", "level": level, "text": line[level:].strip()})
            i += 1
        elif re.match(r'(?:[-+*]|\d+[.)]) ', line):
            ordered = bool(re.match(r'\d', line))
            start = int(re.match(r'\d+', line)[0]) if ordered else 1
            items = []
            pattern = r'\d+[.)] (.*)' if ordered else r'[-+*] (.*)'
            while i < len(lines):
                match = re.fullmatch(pattern, lines[i])
                if not match:
                    break
                item = match[1]; i += 1
                while i < len(lines) and lines[i].startswith('  ') and lines[i].strip():
                    if re.match(r'(?:[-+*]|\d+[.)]) ', lines[i].lstrip()):
                        raise ExportError('nested lists unsupported; use explicit separate lists')
                    item += ' ' + lines[i].strip(); i += 1
                items.append(item)
            output.append({'kind':'list', 'ordered':ordered, 'start':start, 'items':items})
        else:
            content = [line]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                r'(?:#{1,6} |> |```|\s*\||[-+*] |\d+[.)] )', lines[i]
            ):
                content.append(lines[i]); i += 1
            output.append({"kind": "paragraph", "text": ' '.join(content)})
    return output


def bundle(root: Path, documents=None, *, profile='legacy', presentation=None) -> dict:
    if root.is_symlink() or not root.is_dir():
        raise ExportError("bundle root must be a real directory")
    if profile not in ('legacy', 'framed', 'shaped'):
        raise ExportError('unknown presentation profile')
    required = () if profile == 'framed' else DOCUMENTS
    default_documents = (('framed.md',) if (root/'framed.md').exists() else ('pitch.md',)) if profile=='framed' else DOCUMENTS
    documents = tuple(documents) if documents is not None else default_documents
    if not documents or len(set(documents)) != len(documents) or not set(required).issubset(documents):
        raise ExportError("include each required document once")
    if set(documents).intersection(DIAGRAMS):
        raise ExportError('diagrams cannot be prose documents')
    result = {"schema_version": "ShapingDocumentExport.v1", "profile":profile,
              "documents": [], "diagrams": []}
    inventory = {}
    for name in (*documents, *(DIAGRAMS if profile != 'framed' else ())):
        raw = read_source(root, name)
        content = raw.decode("utf-8")
        if not content.strip():
            raise ExportError(f"empty required source: {name}")
        sha = digest(raw)
        inventory[name] = sha
        if name in DIAGRAMS:
            result["diagrams"].append({"id": name, "sha256": sha, "mermaid": content})
        else:
            parsed = blocks(content)
            section, seen = name + '#preamble', set()
            for block in parsed:
                if block['kind'] == 'heading':
                    slug = re.sub(r'[^a-z0-9]+', '-', block['text'].lower()).strip('-') or 'section'
                    unique, suffix = slug, 2
                    while unique in seen:
                        unique = f'{slug}-{suffix}'; suffix += 1
                    seen.add(unique)
                    section = name + '#' + unique
                    block['id'] = section
                block['section_id'] = section
                block['document_id'] = name
            if name in ("contracts.md", "security-owners.md") and not any(
                b["kind"] == "table" and b["rows"] for b in parsed
            ):
                raise ExportError(f"required table has no data rows: {name}")
            result["documents"].append({"id": name, "sha256": sha, "markdown": content, "blocks": parsed})
    result["source_inventory"] = inventory
    result["source_hash"] = digest(json.dumps(inventory, sort_keys=True).encode())
    result["table_row_counts"] = {
        doc["id"]: [len(b["rows"]) for b in doc["blocks"] if b["kind"] == "table"]
        for doc in result["documents"]
    }
    result["verification"] = "source_export_only; semantic, visual and saved-target review required"
    if profile == 'shaped':
        if not presentation or presentation in inventory:
            raise ExportError('shaped profile requires a separate --presentation mapping')
        raw = read_source(root, presentation)
        mapping = json.loads(raw)
        if (not isinstance(mapping, dict) or mapping.get('schema_version') != 'ShapingPresentation.v1'
                or not isinstance(mapping.get('diagrams'), dict)
                or set(mapping['diagrams']) != set(DIAGRAMS)):
            raise ExportError('presentation must map exactly the three diagram IDs')
        anchors = {b['id'] for d in result['documents'] for b in d['blocks'] if b['kind']=='heading'}
        for diagram in result['diagrams']:
            entry = mapping['diagrams'][diagram['id']]
            if not isinstance(entry, dict) or not {'section_id','caption','alt'} <= set(entry) or set(entry)-set(FIGURE_FIELDS):
                raise ExportError('each diagram requires section_id, caption and alt; optional title/legend are supported')
            if any(not isinstance(v, str) or not v.strip() for v in entry.values()):
                raise ExportError('diagram metadata must be nonempty text')
            if entry['section_id'] not in anchors:
                raise ExportError('missing diagram section anchor: '+entry['section_id'])
            diagram.update(entry)
        result['presentation'] = {'path':presentation, 'sha256':digest(raw)}
        table_maps = mapping.get('tables', {})
        if not isinstance(table_maps, dict) or set(table_maps) - {'contracts.md','security-owners.md'}:
            raise ExportError('tables mapping keys must be contracts.md or security-owners.md')
        used_ids = set()
        for doc in result['documents']:
            if doc['id'] in ('contracts.md', 'security-owners.md'):
                tables = [b for b in doc['blocks'] if b['kind']=='table']
                maps = table_maps.get(doc['id'], [None]*len(tables))
                if not isinstance(maps, list) or len(maps) != len(tables):
                    raise ExportError('supply one table map per source table: '+doc['id'])
                for block, table_map in zip(tables, maps):
                    if table_map == {'layout': 'native'}:
                        # Acceptance/example tables retain their own schema. This
                        # is explicit layout selection, never a semantic pass.
                        continue
                    block['presentation'] = compact_table(doc['id'], block, used_ids, table_map)
    elif presentation is not None:
        raise ExportError('--presentation is only supported by the shaped profile')
    return result


def compact_table(document: str, block: dict, used_ids: set, table_map=None) -> dict:
    """A small explicit table vocabulary; detail retains *every* original cell."""
    aliases = {
        'id':'ID', 'interface':'Interface', 'method':'Interface', 'endpoint':'Interface',
        'method endpoint':'Interface', 'method purpose':'Interface',
        'method endpoint purpose':'Interface', 'purpose':'Purpose', 'state':'State',
        'status':'State', 'owner':'Owner', 'owner action':'Owner', 'producer':'detail',
        'consumer':'detail', 'direction':'detail', 'input':'detail', 'inputs':'detail',
        'output':'detail', 'outputs':'detail', 'error':'detail', 'errors':'detail',
        'failure':'detail', 'failures':'detail', 'retry':'detail', 'retries':'detail',
        'error retry':'detail', 'errors retries':'detail', 'consistency':'detail',
        'retry consistency':'detail', 'persistence':'detail', 'security':'detail',
        'evidence':'detail', 'boundary':'detail', 'responsibility':'Responsibility',
        'how enforced':'How enforced', 'enforcement':'How enforced',
        'method interface':'Interface', 'direction producer consumer':'detail',
        'inputs bounds':'detail', 'output meaning':'detail', 'material errors':'detail',
        'timeout retry idempotency':'detail', 'consistency persistence lifecycle':'detail',
        'trust access':'detail', 'state evidence':'State', 'non responsibility':'detail',
        'how enforced control':'How enforced', 'enforcement location':'detail',
        'trust data boundary':'detail', 'evidence state':'detail',
        'explicit negative responsibility and actual owner':'detail', 'gap next action':'detail',
    }
    headers = block['headers']
    normalized = [re.sub(r'[^a-z0-9]+', ' ', h.lower()).strip() for h in headers]
    contract = document == 'contracts.md'
    if len(set(normalized)) != len(headers):
        raise ExportError('duplicate table fields: '+document)
    if table_map is not None:
        if not isinstance(table_map, dict) or set(table_map) != {'id_field','primary','detail'}:
            raise ExportError('table map requires id_field, primary groups and detail fields')
        groups = table_map['primary']
        detail = table_map['detail']
        if (not isinstance(groups, dict) or not 1 <= len(groups) <= 5
                or not isinstance(detail, list) or table_map['id_field'] not in headers):
            raise ExportError('table map requires 1–5 primary groups and a source identity field')
        if any(not isinstance(k, str) or not k.strip() or not isinstance(v, list) or not v for k,v in groups.items()):
            raise ExportError('each primary group must name its source fields')
        fields = [h for v in groups.values() for h in v] + detail
        if any(not isinstance(h, str) for h in fields) or sorted(fields) != sorted(headers):
            raise ExportError('table map must enumerate every source header exactly once')
        if table_map['id_field'] not in next(iter(groups.values())):
            raise ExportError('first primary group must include id_field for linked detail')
        identity = headers.index(table_map['id_field'])
    else:
        if any(h not in aliases for h in normalized):
            raise ExportError('unrecognized table fields in '+document+'; provide --presentation tables mapping with id_field, primary and detail')
        roles = [aliases[h] for h in normalized]
        primary = ['ID','Interface','Purpose','State','Owner'] if contract else ['Responsibility','Owner','How enforced']
        required = ['ID','Interface','State','Owner'] if contract else primary
        if any(h not in roles for h in required):
            raise ExportError('missing primary table fields in '+document+'; provide --presentation tables mapping')
        groups = {role:[h for h,r in zip(headers,roles) if r==role] for role in primary if role in roles}
        identity = roles.index('ID' if 'ID' in roles else 'Responsibility')
    records = []
    for row in block['rows']:
        key = row[identity]
        if not key.strip():
            raise ExportError('empty stable row identity')
        plain_id = bool(re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*', key))
        if contract and not plain_id:
            raise ExportError('contract ID must be a plain stable identifier')
        anchor = document + '--' + (key if contract or (headers[identity].lower()=='id' and plain_id) else digest(key.encode())[:16])
        if anchor in used_ids:
            raise ExportError('duplicate stable row identity: '+key)
        used_ids.add(anchor)
        summary = []
        for fields in groups.values():
            indices = [headers.index(field) for field in fields]
            # Label combined interface/owner fields instead of rewriting their values.
            summary.append(row[indices[0]] if len(indices)==1 else '\n'.join(headers[i]+': '+row[i] for i in indices))
        fields = list(zip(headers, row))
        if [v for _, v in fields] != row or len(fields) != len(headers):
            raise ExportError('table field parity failure')
        records.append({'id':anchor, 'key':key, 'summary':summary, 'fields':fields})
    return {'headers':list(groups), 'records':records}


def safe_link(target: str) -> bool:
    if re.search(r'[\s\x00-\x1f\x7f\\<>"\']', target):
        return False
    try:
        parsed = urlsplit(target)
    except ValueError:
        return False
    return ((parsed.scheme in ('https','http') and bool(parsed.netloc))
            or (parsed.scheme == 'mailto' and bool(parsed.path))
            or (target.startswith('#') and len(target)>1))


def inline(text: str) -> list[dict]:
    """Deliberately bounded inline subset, shared by HTML and native DOCX."""
    result, end = [], 0
    pattern = r'`([^`\n]+)`|\[([^\]\n]+)\]\(([^()\n]+)\)|\*\*([^*\n]+)\*\*|\*([^*\n]+)\*'
    for match in re.finditer(pattern, text):
        if match.start() > end:
            result.append({'text':text[end:match.start()]})
        code, label, target, bold, italic = match.groups()
        if code is not None: result.append({'text':code, 'code':True})
        elif label is not None and (safe_link(target) or relative_link(target)): result.append({'text':label, 'href':target})
        elif label is not None: result.append({'text':match[0]})
        elif bold is not None: result.append({'text':bold, 'bold':True})
        else: result.append({'text':italic, 'italic':True})
        end = match.end()
    if end < len(text): result.append({'text':text[end:]})
    return result


def relative_link(target: str) -> bool:
    return bool(target and not re.search(r'[:\\\x00-\x20\x7f<>"\']',target) and not target.startswith(('/','#')))


def exported_anchors(data: dict) -> set[str]:
    anchors = {d['id'] for d in data['documents']}
    for doc in data['documents']:
        for b in doc['blocks']:
            if b['kind']=='heading': anchors.add(b['id'])
            for record in b.get('presentation',{}).get('records',[]): anchors.add(record['id'])
    return anchors


def resolve_fragment(target: str, document: str, data: dict) -> str | None:
    anchors = exported_anchors(data)
    if relative_link(target):
        path, separator, fragment = target.partition('#')
        destination = posixpath.normpath(posixpath.join(posixpath.dirname(document),path))
        if destination not in {d['id'] for d in data['documents']}: return None
        resolved = destination + ('#'+fragment if separator else '')
        if resolved not in anchors:
            raise ExportError('unresolved source fragment: '+target+' in '+document)
        return '#'+resolved
    if not target.startswith('#'): return target
    fragment = target[1:]
    resolved = fragment if fragment in anchors else document+'#'+fragment
    if resolved not in anchors:
        raise ExportError('unresolved source fragment: '+target+' in '+document)
    return '#'+resolved


def resolved_inline(text: str, document: str, data: dict) -> list[dict]:
    parts = inline(text)
    for part in parts:
        if 'href' in part:
            destination = part['href']
            resolved = resolve_fragment(destination,document,data)
            if resolved is None:
                part['text'] += ' ('+destination+')'; del part['href']
            else: part['href'] = resolved
    return parts


def semantic_text(text: str, document: str, data: dict) -> str:
    return ''.join(p['text'] for p in resolved_inline(text,document,data))


def diagram_metadata_hash(diagram: dict) -> str:
    return digest(json.dumps({k:diagram[k] for k in FIGURE_FIELDS if k in diagram},sort_keys=True,ensure_ascii=False).encode())


def inline_html(text: str, document='', data=None) -> str:
    out = []
    for part in (resolved_inline(text,document,data) if data is not None else inline(text)):
        content = html.escape(part['text'])
        if 'href' in part:
            target = part['href']
            content = '<a href="'+html.escape(target, quote=True)+'">'+content+'</a>'
        for key, tag in [('bold','strong'),('italic','em'),('code','code')]:
            if part.get(key): content = '<'+tag+'>'+content+'</'+tag+'>'
        out.append(content)
    return ''.join(out)


def presentation_blocks(data: dict):
    """Insert figures at the end of their explicit section, before the next heading."""
    figures = {}
    for diagram in data['diagrams']:
        if 'section_id' in diagram:
            figures.setdefault(diagram['section_id'], []).append(diagram)
    for doc in data['documents']:
        yield {'kind':'document','id':doc['id']}
        for i, block in enumerate(doc['blocks']):
            yield block
            following = doc['blocks'][i+1] if i+1<len(doc['blocks']) else None
            if following is None or following['section_id'] != block['section_id']:
                for diagram in figures.get(block['section_id'], []):
                    yield {'kind':'figure', 'diagram':diagram}
    for diagram in data['diagrams']:
        if 'section_id' not in diagram:
            yield {'kind':'figure', 'diagram':diagram}


def html_document(data: dict, images: Path | None = None, *, self_contained=False) -> str:
    if self_contained and data['diagrams'] and images is None:
        raise ExportError('self-contained shaped HTML requires --images with verified SVG/PNG assets')
    verified = verify_assets(data, images, require_svg=True) if images is not None else {}
    native_mermaid = bool(data['diagrams']) and not verified
    esc = html.escape
    title = next((b['text'] for d in data['documents'] for b in d['blocks'] if b['kind']=='heading'), 'Pitch')
    out = ['<!doctype html><html lang="en"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">']
    if not native_mermaid:
        out.append('<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">')
    out.extend(['<title>'+esc(title)+'</title>', '''<style>
body{font:16px/1.55 Arial,sans-serif;max-width:1040px;margin:auto;padding:32px;color:#243344;background:white}
h1,h2{color:#1b5bbb}h1{font-size:24px;line-height:1.2}h2{font-size:20px;margin-top:2em}
h3,h4,h5,h6{color:#157a3b}h3{font-size:18px;margin-top:1.6em}a{color:#184e9d}
p,ul,ol{margin:.7em 0 1em}li{margin:.3em 0}table{border-collapse:collapse;width:100%;min-width:560px}
th,td{border:1px solid #d9d9d9;padding:9px 11px;text-align:left;vertical-align:top;white-space:pre-line;word-break:normal}
th{background:#183b56;color:white}tbody tr:nth-child(even){background:#f3f6fa}
thead{display:table-header-group}.table,.mermaid,.full-size{overflow:auto}.table{margin:1em 0}
.table:focus-visible,a:focus-visible,summary:focus-visible{outline:3px solid #157a3b;outline-offset:3px}
pre{white-space:pre-wrap;overflow-wrap:break-word}code{overflow-wrap:anywhere}figure{margin:1.7em 0}figure img{max-width:100%;height:auto}
figcaption{font-style:italic;color:#495565;margin:.6em 0}details{margin:1em 0}summary{cursor:pointer;color:#184e9d}
dt{font-weight:bold;color:#157a3b;margin-top:.6em}dd{margin:.2em 0 .8em;white-space:pre-wrap;overflow-wrap:break-word}
.figure-view{overflow:auto}.figure-toggle:checked~.figure-view img{max-width:none}
.figure-toggle:focus-visible+label{outline:3px solid #157a3b}.figure-legend{white-space:pre-wrap}
blockquote{margin:1em 1.5em;font-style:italic}.source{font-size:12px;overflow-wrap:break-word;color:#495565}
@media(max-width:600px){body{padding:16px}h1{font-size:24px}.table{max-width:100%}}
@page{size:letter;margin:.65in}@media print{body{font-size:11pt;padding:0}h1{font-size:18pt}h2{font-size:13pt}h3{font-size:11.5pt}table{min-width:0}h1,h2,h3,figcaption{break-after:avoid}.table{overflow:visible}.figure-toggle,.figure-toggle+label{display:none}.figure-toggle:checked~.figure-view img{max-width:100%}figure{break-inside:avoid}}
</style></head><body><main>'''])
    for b in presentation_blocks(data):
        kind = b['kind']
        def formatted(text):
            return inline_html(text,b.get('document_id',''),data)
        if kind == 'document':
            out.append('<span id="'+esc(b['id'],quote=True)+'"></span>')
        elif kind == 'heading':
            n = b['level']
            out.append(f'<h{n} id="{esc(b["id"], quote=True)}">{formatted(b["text"])}</h{n}>')
        elif kind == 'table':
            compact = b.get('presentation')
            headers = compact['headers'] if compact else b['headers']
            rows = [r['summary'] for r in compact['records']] if compact else b['rows']
            out.append('<div class="table" tabindex="0" role="region" aria-label="Scrollable table"><table><thead><tr>' + ''.join('<th>'+esc(c)+'</th>' for c in headers) + '</tr></thead><tbody>')
            for index, row in enumerate(rows):
                out.append('<tr>')
                for column, cell in enumerate(row):
                    content = formatted(cell)
                    if compact and column == 0:
                        anchor = esc(compact['records'][index]['id'], quote=True)
                        content = '<a href="#'+anchor+'">'+content+'</a>'
                    out.append('<td>'+content+'</td>')
                out.append('</tr>')
            out.append('</tbody></table></div>')
            if compact:
                for record in compact['records']:
                    # Open detail remains printable and directly reachable by fragment.
                    out.append('<details open id="'+esc(record['id'], quote=True)+'"><summary>'+esc(record['key'])+' — all source fields</summary><dl>')
                    for field, value in record['fields']:
                        out.append('<dt>'+esc(field)+'</dt><dd>'+formatted(value)+'</dd>')
                    out.append('</dl></details>')
        elif kind == 'code':
            out.append('<pre>'+esc(b['text'])+'</pre>')
        elif kind == 'list':
            tag = 'ol' if b['ordered'] else 'ul'
            out.append('<'+tag+(f' start="{b["start"]}"' if b['ordered'] else '')+'>')
            out.extend('<li>'+formatted(item)+'</li>' for item in b['items'])
            out.append('</'+tag+'>')
        elif kind == 'figure':
            d = b['diagram']; caption = d.get('caption',d['id'])
            out.append('<figure id="figure-'+esc(d['id'],quote=True)+'">')
            if d.get('title'): out.append('<h3>'+esc(d['title'])+'</h3>')
            if verified:
                asset = verified[d['id']]
                uri = 'data:image/svg+xml;base64,'+base64.b64encode(asset['svg']).decode()
                width, height = asset['svg_dimensions']
                img = '<img src="'+uri+'" width="'+str(width)+'" height="'+str(height)+'" alt="'+esc(d.get('alt',caption),quote=True)+'">'
                control = 'expand-'+esc(d['id'],quote=True)
                out.append('<input class="figure-toggle" type="checkbox" id="'+control+'"><label for="'+control+'">Full-size figure</label>')
                out.append('<div class="figure-view">'+img+'</div>')
                out.append('<figcaption>'+esc(caption)+'</figcaption>')
            else:
                out.append('<pre class="mermaid">'+esc(d['mermaid'])+'</pre>')
                out.append('<figcaption>'+esc(caption)+'</figcaption>')
            if d.get('legend'): out.append('<p class="figure-legend">'+esc(d['legend'])+'</p>')
            out.append('</figure>')
        elif kind == 'quote':
            out.append('<blockquote><p>'+formatted(b['text'])+'</p></blockquote>')
        else:
            out.append('<p>'+formatted(b['text'])+'</p>')
    out.append('<p class="source">Source revision '+data['source_hash']+'</p>')
    if native_mermaid:
        out.append('<script type="module">import mermaid from "https://cdn.jsdelivr.net/npm/mermaid@11.17.2/dist/mermaid.esm.min.mjs";mermaid.initialize({startOnLoad:false,securityLevel:"strict",htmlLabels:false});await mermaid.run({querySelector:".mermaid"});</script>')
    out.append('</main></body></html>')
    return '\n'.join(out)


@contextmanager
def directory_handle(path: Path, *, create=True):
    """Pin each physical output directory; never traverse a replaced symlink.

    POSIX dir-fd/no-follow primitives are required rather than silently falling
    back to a racy pathname publication on unsupported hosts.
    """
    if os.name != 'posix' or not hasattr(os, 'O_NOFOLLOW'):
        raise ExportError('safe output publication requires POSIX directory descriptors')
    path = Path(os.path.abspath(path))
    fd = os.open(path.anchor, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        for part in path.parts[1:]:
            try:
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=fd)
            except FileNotFoundError:
                if not create:
                    raise
                try:
                    os.mkdir(part, mode=0o700, dir_fd=fd)
                except FileExistsError:
                    pass
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=fd)
            os.close(fd); fd = child
        yield fd
    except OSError as error:
        raise ExportError(f'unsafe or unavailable output directory: {error}') from error
    finally:
        os.close(fd)


def write_at(directory: int, name: str, raw: bytes, identical=True) -> None:
    if not name or Path(name).name != name:
        raise ExportError('invalid output filename')
    try:
        fd = os.open(name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                     0o600, dir_fd=directory)
    except FileExistsError:
        if identical:
            try:
                fd = os.open(name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory)
                with os.fdopen(fd, 'rb') as stream:
                    if stat.S_ISREG(os.fstat(stream.fileno()).st_mode) and stream.read(len(raw)+1) == raw:
                        return
            except OSError:
                pass
        raise ExportError(f'output exists; use a new revision directory: {name}')
    with os.fdopen(fd, 'wb') as stream:
        stream.write(raw); stream.flush(); os.fsync(stream.fileno())
    os.fsync(directory)


def write_new(path: Path, raw: bytes) -> None:
    with directory_handle(path.parent) as directory:
        write_at(directory, path.name, raw)


def run_renderer(argv: list[str], timeout: float) -> str:
    """Finite wall-clock limit, bounded diagnostics, owned process-group cleanup."""
    if timeout <= 0 or timeout > 300 or os.name != 'posix':
        raise ExportError('renderer needs POSIX and a timeout in (0, 300] seconds')
    diagnostic = bytearray()
    process = subprocess.Popen(argv, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                               stdin=subprocess.DEVNULL, start_new_session=True)
    deadline = time.monotonic() + timeout
    timed_out = False
    try:
        with selectors.DefaultSelector() as selector:
            selector.register(process.stdout, selectors.EVENT_READ)
            while selector.get_map():
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    timed_out = True
                    break
                for key, _ in selector.select(min(remaining, .1)):
                    chunk = os.read(key.fileobj.fileno(), 8192)
                    if not chunk:
                        selector.unregister(key.fileobj)
                    elif len(diagnostic) < 16384:
                        diagnostic.extend(chunk[:16384-len(diagnostic)])
            if not timed_out:
                try:
                    process.wait(timeout=max(.001, deadline-time.monotonic()))
                except subprocess.TimeoutExpired:
                    timed_out = True
    finally:
        # Only this explicitly started group, never a broad renderer/process match.
        for sig in (signal.SIGTERM, signal.SIGKILL):
            try:
                os.killpg(process.pid, sig)
            except ProcessLookupError:
                pass
            if sig == signal.SIGTERM:
                try:
                    process.wait(timeout=.5)
                except subprocess.TimeoutExpired:
                    pass
        process.wait()
        process.stdout.close()
    if timed_out:
        raise ExportError(f'renderer timed out after {timeout} seconds')
    if process.returncode:
        raise ExportError(f'renderer exit {process.returncode}: {diagnostic.decode("utf-8",errors="replace")}')
    return diagnostic.decode('utf-8', errors='replace').strip()


def parse_svg(raw: bytes):
    # Decode before declaration checks; never let XML autodetect UTF-16/entities.
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeError as error:
        raise ExportError('SVG must be UTF-8') from error
    if len(raw) > 20_000_000 or '\x00' in text or re.search(r'<!DOCTYPE|<!ENTITY|<\?', text, re.I):
        raise ExportError('unsafe SVG declaration or oversized asset')
    try:
        return ET.fromstring(text)
    except ET.ParseError as error:
        raise ExportError('invalid SVG XML') from error


def passive_css(value):
    if re.search(r'[\\@<]|/\*|expression\s*\(|(?:https?|file|data|javascript)\s*:|(?:-webkit-)?image(?:-set)?\s*\(|\banimation(?:-[\w-]+)?\s*:', value, re.I):
        raise ExportError('unsafe SVG CSS/resource or animation reference')
    without_local = re.sub(r'url\(\s*[\'"]?#[A-Za-z_][\w:.-]*[\'"]?\s*\)', '', value, flags=re.I)
    if re.search(r'url\s*\(', without_local, re.I):
        raise ExportError('external SVG resource refused')


def normalize_renderer_svg(raw: bytes) -> tuple[bytes, list[str]]:
    """Remove only Mermaid's exact built-in animation CSS when it is unused.

    Called on renderer output, never on source or imported assets. Active/unknown
    CSS is rejected by the full validator after these recorded transformations.
    """
    root = parse_svg(raw)
    forbidden_classes = {'edge-animation-slow','edge-animation-fast'}
    if any(forbidden_classes.intersection(el.get('class','').split()) for el in root.iter()):
        raise ExportError('source-requested Mermaid animation is unsupported')
    rules = [
        '@keyframes edge-animation-frame{from{stroke-dashoffset:0;}}',
        '@keyframes dash{to{stroke-dashoffset:0;}}',
    ]
    identifier = root.get('id','')
    for speed, seconds in [('slow',50),('fast',20)]:
        rules.append(f'#{identifier} .edge-animation-{speed}'+'{stroke-dasharray:9,5!important;stroke-dashoffset:900;'+f'animation:dash {seconds}s linear infinite;stroke-linecap:round;'+'}')
    transforms = []
    for style in root.iter('{http://www.w3.org/2000/svg}style'):
        if len(style): raise ExportError('SVG style must not contain child elements')
        css = style.text or ''
        for rule in rules:
            if rule in css:
                css = css.replace(rule,'')
                transforms.append('removed unused built-in CSS: '+rule)
        style.text = css
    normalized = ET.tostring(root,encoding='utf-8') if transforms else raw
    svg_dimensions(normalized)
    return normalized, transforms


def validate_mermaid(source: str) -> dict:
    """Preserve static JSON init styling verbatim; refuse resource/security overrides."""
    allowed, settings = {'theme','themeVariables','themeCSS'}, {}
    pattern = r'%%\{\s*init\s*:\s*(\{.*?\})\s*\}%%'
    for match in re.finditer(pattern,source,re.S):
        try:
            config = json.loads(match[1])
        except json.JSONDecodeError as error:
            raise ExportError('Mermaid init styling requires strict JSON') from error
        if not isinstance(config,dict) or set(config)-allowed:
            raise ExportError('Mermaid init supports only static theme/themeVariables/themeCSS; security overrides refused')
        if 'theme' in config and config['theme'] not in ('default','base','dark','forest','neutral','null'):
            raise ExportError('unsupported Mermaid theme')
        variables = config.get('themeVariables',{})
        if not isinstance(variables,dict): raise ExportError('themeVariables must be a flat object')
        for key,value in variables.items():
            if not re.fullmatch(r'[A-Za-z][A-Za-z0-9]*',key) or not isinstance(value,(str,int,float,bool)):
                raise ExportError('unsupported Mermaid theme variable')
            passive_css(str(value))
        css = config.get('themeCSS','')
        if not isinstance(css,str): raise ExportError('themeCSS must be static text')
        passive_css(css)
        settings.update(config)
    remainder = re.sub(pattern,'',source,flags=re.S)
    if re.search(r'%%\s*\{|^---|(?:https?|file|data|javascript)\s*:|\bclick\s|\b(?:animate|img)\s*:|\b(?:url|image-set)\s*\(',remainder,re.I|re.M):
        raise ExportError('unsupported Mermaid directives, animation or resource references')
    return settings


def svg_dimensions(raw: bytes) -> list[float]:
    """Validate passive UTF-8 SVG. Imported assets are never normalized here."""
    root = parse_svg(raw)
    namespace = '{http://www.w3.org/2000/svg}'
    tags = {'svg','g','defs','marker','path','rect','line','polyline','polygon','circle',
            'ellipse','text','tspan','title','desc','style','clipPath','linearGradient',
            'radialGradient','stop','use','symbol','filter','feDropShadow'}

    ids, references = set(), []
    for element in root.iter():
        if not isinstance(element.tag, str) or not element.tag.startswith(namespace) or element.tag[len(namespace):] not in tags:
            raise ExportError('active or unsupported SVG element: '+str(element.tag))
        tag = element.tag[len(namespace):]
        if tag == 'style':
            if len(element): raise ExportError('SVG style must not contain child elements')
            passive_css(element.text or '')
        for attribute, value in element.attrib.items():
            if attribute.startswith('{') and attribute != '{http://www.w3.org/1999/xlink}href':
                if attribute != '{http://www.w3.org/XML/1998/namespace}space':
                    raise ExportError('unsupported SVG attribute namespace')
            local = attribute.rsplit('}',1)[-1].lower()
            if local.startswith('on') or local in ('src','srcset','base','tabindex'):
                raise ExportError('active SVG attribute')
            if local == 'href':
                if not re.fullmatch(r'#[A-Za-z_][\w:.-]*',value):
                    raise ExportError('external SVG reference refused')
                references.append(value[1:])
            elif local == 'id':
                if value in ids: raise ExportError('duplicate SVG ID')
                ids.add(value)
            else:
                passive_css(value)
    if root.tag != namespace+'svg' or any(ref not in ids for ref in references):
        raise ExportError('invalid SVG root or missing local reference')
    try:
        if 'viewBox' in root.attrib:
            box = [float(n) for n in re.split(r'[\s,]+', root.attrib['viewBox'].strip())]
            if len(box) != 4 or not all(math.isfinite(n) for n in box): raise ValueError()
            dimensions = box[2:]
        else:
            dimensions = [float(re.sub(r'px$', '', root.attrib[k])) for k in ('width','height')]
        if not all(math.isfinite(n) and 0 < n <= 100000 for n in dimensions): raise ValueError()
    except (KeyError,ValueError):
        raise ExportError('SVG needs finite positive viewBox or pixel dimensions') from None
    return dimensions


def png_dimensions(raw: bytes) -> list[int]:
    if not raw.startswith(b'\x89PNG\r\n\x1a\n') or len(raw)>80_000_000:
        raise ExportError('invalid or oversized PNG')
    position, tags, dimensions = 8, [], None
    while position+12 <= len(raw):
        length = struct.unpack('>I',raw[position:position+4])[0]
        end = position+12+length
        if end > len(raw): raise ExportError('truncated PNG')
        tag, content = raw[position+4:position+8], raw[position+8:end-4]
        if zlib.crc32(tag+content) != struct.unpack('>I',raw[end-4:end])[0]:
            raise ExportError('PNG checksum mismatch')
        tags.append(tag)
        if tag == b'IHDR':
            if length != 13 or len(tags)!=1: raise ExportError('invalid PNG header')
            dimensions = list(struct.unpack('>II',content[:8]))
            if not all(0<n<=100000 for n in dimensions): raise ExportError('invalid PNG dimensions')
        position = end
        if tag == b'IEND': break
    if position != len(raw) or not dimensions or b'IDAT' not in tags or tags[-1] != b'IEND':
        raise ExportError('incomplete PNG')
    return dimensions


def verify_assets(data: dict, images: Path, *, require_svg=False) -> dict:
    receipts = json.loads(read_source(images, 'render-manifest.json'))
    if not isinstance(receipts, dict): raise ExportError('invalid render manifest')
    verified = {}
    for diagram in data['diagrams']:
        entry = receipts.get(diagram['id'], {})
        if not isinstance(entry, dict) or entry.get('source_sha256') != diagram['sha256']:
            raise ExportError('stale or missing render source binding')
        for key in ('title','legend'):
            if entry.get(key) != diagram.get(key):
                raise ExportError('stale render presentation binding: '+key)
        for key in FIGURE_FIELDS:
            if key in diagram and entry.get(key) != diagram[key]:
                raise ExportError('stale render presentation binding: '+key)
        if ('presentation_sha256' in entry or any(k in diagram for k in ('title','legend'))):
            if entry.get('presentation_sha256') != diagram_metadata_hash(diagram):
                raise ExportError('stale render presentation metadata hash')
        rich = require_svg or data.get('profile') == 'shaped'
        if rich:
            renderer = entry.get('renderer', {})
            if (not isinstance(renderer, dict) or not isinstance(renderer.get('version'),str)
                    or not renderer['version'].strip() or not isinstance(renderer.get('config'),dict)
                    or renderer['config'].get('htmlLabels') is not False
                    or renderer['config'].get('securityLevel') != 'strict'):
                raise ExportError('missing safe renderer identity/configuration')
        item = {}
        for extension, inspect in [('png',png_dimensions), *([('svg',svg_dimensions)] if rich else [])]:
            raw = read_source(images, entry.get(extension, ''))
            if digest(raw) != entry.get(extension+'_sha256'):
                raise ExportError('render image hash mismatch: '+extension)
            dimensions = inspect(raw)
            if rich and entry.get(extension+'_dimensions') != dimensions:
                raise ExportError('render image dimensions mismatch: '+extension)
            item[extension], item[extension+'_dimensions'] = raw, dimensions
        verified[diagram['id']] = item
    return verified


def render(data: dict, root: Path, output: Path, executable: str, timeout=60, *, node=None, puppeteer_config=None) -> dict:
    # Refuse symlink parents before creating a directory or invoking a renderer.
    with directory_handle(output.parent) as parent:
        try:
            os.mkdir(output.name, mode=0o700, dir_fd=parent)
        except FileExistsError as error:
            raise ExportError('render to a fresh directory') from error
        directory = os.open(output.name, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=parent)
        try:
            receipts = {}
            with tempfile.TemporaryDirectory(prefix='engos-mermaid-') as temporary:
                scratch = Path(temporary).resolve()
                command = [node,executable] if node else [executable]
                version = run_renderer([*command, '--version'], timeout)
                if not version:
                    raise ExportError('renderer did not report its version')
                config = scratch/'mermaid-config.json'
                config.write_text(json.dumps(RENDER_CONFIG), encoding='utf-8')
                browser_args, browser_sha = [], None
                if puppeteer_config is not None:
                    raw_config = read_source(puppeteer_config.parent,puppeteer_config.name)
                    browser = scratch/'puppeteer.json'; browser.write_bytes(raw_config)
                    browser_args = ['-p',str(browser)]; browser_sha = digest(raw_config)
                for diagram in data['diagrams']:
                    styling = validate_mermaid(diagram['mermaid'])
                    source = scratch / diagram['id']
                    source.write_text(diagram['mermaid'], encoding='utf-8')
                    entry = {'source_sha256':diagram['sha256'], 'pixel_inspection':'pending',
                             'renderer':{'executable':executable, 'version':version,
                                         'node':node, 'puppeteer_config_sha256':browser_sha,
                                         'source_styling':styling, 'config':RENDER_CONFIG,
                                         'width':1400, 'scale':2, 'background':'white'}}
                    entry.update({k:diagram.get(k, diagram['id'] if k!='section_id' else None) for k in ('section_id','caption','alt')})
                    entry.update({k:diagram[k] for k in ('title','legend') if k in diagram})
                    entry['presentation_sha256'] = diagram_metadata_hash(diagram)
                    for extension, inspect in [('svg',svg_dimensions),('png',png_dimensions)]:
                        target = scratch / (source.stem + '.' + extension)
                        run_renderer([*command, '-i', str(source), '-o', str(target), '-c', str(config),
                                      *browser_args, '-b', 'white', '-w', '1400', '-s', '2'], timeout)
                        raw = read_source(scratch, target.name)
                        if extension == 'svg':
                            entry['svg_original_sha256'] = digest(raw)
                            raw, entry['svg_normalization'] = normalize_renderer_svg(raw)
                        dimensions = inspect(raw)
                        write_at(directory, target.name, raw, identical=False)
                        entry.update({extension:target.name, extension+'_sha256':digest(raw),
                                      extension+'_dimensions':dimensions})
                    if digest(read_source(root, diagram['id'])) != diagram['sha256']:
                        raise ExportError('source changed during rendering')
                    receipts[diagram['id']] = entry
                for name, sha in data['source_inventory'].items():
                    if digest(read_source(root,name)) != sha:
                        raise ExportError('source changed during rendering')
                if 'presentation' in data and digest(read_source(root,data['presentation']['path'])) != data['presentation']['sha256']:
                    raise ExportError('presentation changed during rendering')
            write_at(directory, 'render-manifest.json', json.dumps(receipts, indent=2).encode(), identical=False)
            return receipts
        finally:
            os.close(directory)


def docx_document(data: dict, images: Path | None, output: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.enum.section import WD_SECTION_START

    if output.exists() or any(p.is_symlink() for p in (output, *output.parents)):
        raise ExportError("DOCX output exists; use a fresh revision")
    # Check every image/source binding BEFORE constructing a document.
    if data['diagrams'] and images is None:
        raise ExportError('docx requires --images with source-bound render-manifest.json')
    verified_images = verify_assets(data, images) if data['diagrams'] else {}
    doc = Document()
    section = doc.sections[0]
    legacy = data.get('profile', 'legacy') == 'legacy'
    section.page_width, section.page_height = (Inches(11.7), Inches(8.3)) if legacy else (Inches(8.5), Inches(11))
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    for name, size in [('Normal',11),('Title',18),('Heading 1',13),('Heading 2',11.5)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor.from_string(GREEN if name=='Heading 2' else BLUE if name!='Normal' else '243344')
        style.font.bold = name != 'Normal'
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.space_before = Pt(18 if name.startswith('Heading') else 0)
        # Explicitly suppress template/theme title border residue.
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    title_added = False
    if not legacy:
        # Do not inherit Word's accent1 caption color; small text needs the same
        # explicit high-contrast neutral used by the HTML presentation.
        doc.styles['Caption'].font.color.rgb = RGBColor.from_string('495565')
    bookmarks = {}
    current_document = ''
    pending_document = None

    def bookmark(p, identity):
        name = 'ref_'+digest(identity.encode())[:28]
        if name in bookmarks: raise ExportError('duplicate DOCX bookmark')
        number = str(len(bookmarks)+1); bookmarks[name] = identity
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'),number); start.set(qn('w:name'),name)
        end.set(qn('w:id'),number)
        p._p.append(start); p._p.append(end)
        return name

    def link(p, text, target, internal=False):
        el = OxmlElement('w:hyperlink')
        if internal: el.set(qn('w:anchor'),'ref_'+digest(target.encode())[:28])
        else: el.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
        run, props, color, content = (OxmlElement(n) for n in ('w:r','w:rPr','w:color','w:t'))
        color.set(qn('w:val'),BLUE); props.append(color); run.append(props)
        content.text = text; run.append(content); el.append(run); p._p.append(el)

    def add_inline(p, text):
        for part in resolved_inline(text,current_document,data):
            if 'href' in part:
                target = part['href']
                link(p,part['text'],target[1:] if target.startswith('#') else target, target.startswith('#'))
            else:
                run = p.add_run(part['text'])
                if part.get('bold'): run.bold = True
                if part.get('italic'): run.italic = True
                if part.get('code'): run.font.name = 'Courier New'

    def bind_document(p):
        nonlocal pending_document
        if pending_document:
            bookmark(p,pending_document); pending_document = None

    def paragraph(text, style=None, *, literal=False):
        p = doc.add_paragraph(style=style)
        bind_document(p)
        p.paragraph_format.widow_control = True
        if literal:
            p.add_run(text)
        else:
            add_inline(p,text)
        return p

    available_width = (section.page_width-section.left_margin-section.right_margin)/914400

    def table(headers, rows, records=None, detail=False):
        native = doc.add_table(rows=0,cols=len(headers)); native.style = 'Table Grid'; native.autofit = False
        weights = [1,3] if detail else [0.65 if h=='ID' else 1.5 if h in ('Interface','Purpose','How enforced','Value') else 1 for h in headers]
        widths = [Inches(available_width*w/sum(weights)) for w in weights]
        page_points = (section.page_height-section.top_margin-section.bottom_margin)/12700
        # Conservative 11pt character budget, including explicit newlines and padding.
        # This is a preflight bound, not measured font pagination; pixels remain a gate.
        heights = []
        for values in [headers,*rows]:
            lines = []
            for value,width in zip(values,widths):
                text = semantic_text(value,current_document,data)
                columns = max(1,int(width.pt/11)-2)
                lines.append(sum(max(1,math.ceil(len(line)/columns)) for line in text.split('\n')))
            heights.append(max(lines)*15+10)
        if any(h+heights[0]+24 > page_points for h in heights[1:]):
            raise ExportError('table row exceeds estimated page capacity; split the source row deliberately: '+headers[-1])
        if detail and sum(heights[:3])+42 > page_points:
            raise ExportError('detail heading/header/first rows exceed estimated page capacity: '+headers[-1])
        for column, width in zip(native.columns,widths): column.width = width
        for index, values in enumerate([headers,*rows]):
            row = native.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for column, (cell,value,width) in enumerate(zip(row.cells,values,widths)):
                cell.width = width
                p = cell.paragraphs[0]
                bind_document(p)
                if records and index and column == 0:
                    link(p,value,records[index-1]['id'],True)
                elif index == 0: p.add_run(value)
                else: add_inline(p,value)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_with_next = index == 0 or (detail and index == 1)
                for run in p.runs:
                    run.font.size = Pt(11)
                    if index == 0:
                        run.bold = True; run.font.color.rgb = RGBColor.from_string('FFFFFF')
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'183B56' if index==0 else 'F3F6FA' if index%2==0 else 'FFFFFF'); props.append(shade)
                borders = OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el = OxmlElement('w:'+edge)
                    for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+key),val)
                    borders.append(el)
                props.append(borders)
        doc.add_paragraph()

    def numbered_list(block):
        numbering = doc.part.numbering_part.element
        abstract_id = max([int(el.get(qn('w:abstractNumId'))) for el in numbering.findall(qn('w:abstractNum'))]+[-1])+1
        abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),str(abstract_id))
        level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'),'0')
        for tag, value in [('start',str(block['start'])),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
            el = OxmlElement('w:'+tag); el.set(qn('w:val'),value); level.append(el)
        abstract.append(level); numbering.append(abstract)
        number = numbering.add_num(abstract_id)
        for item in block['items']:
            p = paragraph(item,'List Number')
            props = p._p.get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val = 0; props.get_or_add_numId().val = number.numId

    for block in presentation_blocks(data):
        kind = block['kind']
        current_document = block.get('document_id',current_document)
        if kind == 'document':
            pending_document = block['id']; current_document = block['id']
        elif kind == 'heading':
            style = 'Title' if not title_added else ('Heading 1' if block['level'] <= 2 else 'Heading 2')
            bookmark(paragraph(block['text'],style),block['id']); title_added = True
        elif kind == 'table':
            compact = block.get('presentation')
            if compact:
                table(compact['headers'],[r['summary'] for r in compact['records']],compact['records'])
                for record in compact['records']:
                    heading = paragraph(record['key']+' — all source fields','Heading 2')
                    heading.paragraph_format.keep_with_next = True
                    bookmark(heading,record['id'])
                    table(['Field',record['key']],record['fields'],detail=True)
            else: table(block['headers'],block['rows'])
        elif kind == 'code':
            paragraph('').add_run(block['text']).font.name = 'Courier New'
        elif kind == 'quote':
            paragraph(block['text'],'Quote')
        elif kind == 'list':
            if block['ordered']: numbered_list(block)
            else:
                for item in block['items']: paragraph(item,'List Bullet')
        elif kind == 'figure':
            diagram = block['diagram']; asset = verified_images[diagram['id']]
            width, height = asset['png_dimensions']
            if legacy and (height > .75*width) != (section.page_height > section.page_width):
                section = doc.add_section(WD_SECTION_START.NEW_PAGE)
                portrait = height > .75*width
                section.page_width, section.page_height = (Inches(8.3),Inches(11.7)) if portrait else (Inches(11.7),Inches(8.3))
            available_width = (section.page_width-section.left_margin-section.right_margin)/914400
            usable_height = (section.page_height-section.top_margin-section.bottom_margin)/914400
            caption = diagram.get('caption',diagram['id'])
            title, legend = diagram.get('title'), diagram.get('legend')
            # Include every supplied figure-text paragraph; no legends inferred from color.
            text_height = sum(max(.35,(sum(max(1,math.ceil(len(line)/max(1,int(available_width*8)))) for line in text.split('\n'))+1)*.24) for text in (title,caption,legend) if text)
            reserved_height = max(.5,text_height+.2)
            if usable_height-reserved_height < 1:
                raise ExportError('figure title/caption/legend exceeds page capacity')
            if title:
                p = paragraph(title,'Heading 2',literal=True)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
            # Keep caption and figure together; pixel readability remains a visual gate.
            p = paragraph(caption,'Caption',literal=True)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            picture = paragraph('')
            picture.paragraph_format.keep_with_next = bool(legend)
            picture.paragraph_format.keep_together = True
            shape = picture.add_run().add_picture(io.BytesIO(asset['png']),width=Inches(min(available_width,(usable_height-reserved_height)*width/height)))
            shape._inline.docPr.set('descr',diagram.get('alt',diagram['id']))
            shape._inline.docPr.set('title',title or caption)
            if legend:
                p = paragraph(legend,literal=True)
                p.paragraph_format.keep_together = True
                p.paragraph_format.keep_with_next = False
        else: paragraph(block['text'])
    doc.add_paragraph('Source revision '+data['source_hash'])
    stream = io.BytesIO()
    doc.save(stream)
    with directory_handle(output.parent) as parent:
        write_at(parent, output.name, stream.getvalue(), identical=False)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('action', choices=('json','html','render','docx'))
    p.add_argument('--bundle', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    p.add_argument('--documents', nargs='+', help='explicit prose inventory; framed prefers framed.md (pitch.md fallback), legacy/shaped require the three default documents')
    p.add_argument('--profile', choices=('legacy','framed','shaped'), default='legacy')
    p.add_argument('--presentation', help='bundle-relative JSON mapping for shaped sections and optional tables')
    p.add_argument('--self-contained', action='store_true', help='HTML only: require local verified assets; no runtime CDN')
    p.add_argument('--mmdc', default='mmdc', help='existing trusted Mermaid CLI executable; no installation performed')
    p.add_argument('--node', help='optional approved Node executable for a cached Mermaid CLI JavaScript entrypoint')
    p.add_argument('--puppeteer-config', type=Path, help='operator-selected local browser config; snapshotted and hash recorded')
    p.add_argument('--images', type=Path)
    p.add_argument('--timeout-seconds', type=float, default=60, help='renderer wall-clock limit per diagram, maximum 300')
    args = p.parse_args(argv)
    try:
        if args.self_contained and args.action != 'html':
            raise ExportError('--self-contained applies only to html')
        data = bundle(args.bundle, args.documents, profile=args.profile, presentation=args.presentation)
        if args.action == 'json':
            write_new(args.output, (json.dumps(data, indent=2, ensure_ascii=False)+'\n').encode())
        elif args.action == 'html':
            write_new(args.output, html_document(data,args.images,self_contained=args.self_contained).encode())
        elif args.action == 'render':
            render(data, args.bundle, args.output, args.mmdc, args.timeout_seconds,node=args.node,puppeteer_config=args.puppeteer_config)
        else:
            docx_document(data, args.images, args.output)
        print(json.dumps({'status':'created_or_identical','source_hash':data['source_hash'],
                          'table_row_counts':data['table_row_counts'], 'diagrams':len(data['diagrams']),
                          'visual_review':'pending','saved_target_review':'not_performed'}))
        return 0
    except (ExportError, OSError, UnicodeError, json.JSONDecodeError, subprocess.CalledProcessError, ImportError) as error:
        print(json.dumps({'status':'error','detail':str(error)}), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
